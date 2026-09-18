"""
skills/excel_skill.py – Read and edit Excel workbooks through the real API
===========================================================================
Excel is the one desktop app NEBULA already had the wrong tools for. There is
UI-automation machinery in ``skills/desktop_skill.py`` and it must never be
pointed at a spreadsheet: keystrokes into a grid are unverifiable, silently
land in the wrong cell when focus moves, and cannot be undone. Excel exposes a
real API through two libraries, so that is the layer this skill works at.

The design is two tools split by risk, exactly like the browser pair:

  * ``excel_read`` -- sheets, shape, read_range. Nothing can be destroyed, so
    it is unconfirmed, and the model is told to call it *first*. Guessing that
    the totals live in column D and writing there is the failure this prevents.
  * ``excel_write`` -- set_values, set_formula, add_sheet, format_bold, behind
    ``confirm=True``, a backup and an error scan.

Three rules carry most of the safety:

  * **the op set is fixed.** An unrecognised op is refused with one sentence
    rather than approximated. A model that has invented "delete_rows" must be
    told no, not met halfway;
  * **a backup before the first write of a session to each workbook.** Ctrl+Z
    does not cross the COM boundary and openpyxl rewrites the whole file, so
    the copy in ``data/excel_backups`` is the entire undo story;
  * **an error-value scan after every write.** ``=B2*0.18`` against a text
    column produces ``#VALUE!`` and Excel reports it nowhere -- the file just
    quietly holds a broken column. Re-reading the written range turns that
    into tool-result text the agent loop can feed back to the model. There is
    deliberately **no retry loop here**: the loop already retries on feedback,
    and a skill that retried privately would burn the backup on its own
    guesses.

The two backends are not equivalent and the difference matters. openpyxl edits
the file on disk and cannot recalculate anything, so a formula it writes has
no value at all until Excel next opens the file -- every openpyxl write says
so in its reply rather than implying a check it could not perform. xlwings
drives the live application, so its scan reads what Excel actually computed.
That is why the live path is preferred whenever Excel already has the file.
"""

from __future__ import annotations

import json
import os
import re
import shutil
from copy import copy as _copy_style
from datetime import date, datetime, time
from pathlib import Path
from typing import TYPE_CHECKING, Any

from config.logging_config import get_logger
from skills.base import Skill

if TYPE_CHECKING:
    from intelligence.task import Task

logger = get_logger("skills.excel")

#: The values Excel puts in a cell when a formula fails. A write that produces
#: one of these looks like a success from the outside -- the save worked, the
#: cell has content -- which is exactly why it has to be looked for.
ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!"}

#: The one sentence an unmapped op gets. Fixed wording, no improvisation.
REFUSAL = (
    "I can read ranges, set values or formulas, add sheets and bold headers "
    "- not that."
)

READ_OPS = ("sheets", "shape", "read_range")
WRITE_OPS = ("set_values", "set_formula", "add_sheet", "format_bold")

#: Data rows a ``shape`` call returns. The point of shape is to let the model
#: see the layout before it writes; five rows is enough to tell a header from
#: data and a date column from a text one, and a 10,000-row sheet must never
#: reach an 8k context.
SHAPE_ROWS = 5

#: Rows sampled -- not returned -- when guessing what a column holds.
TYPE_SAMPLE_ROWS = 20

#: Cells ``read_range`` will read back before it truncates. A model asking for
#: A1:XFD1048576 gets the top-left corner and a note, not the workbook.
MAX_RANGE_CELLS = 200

#: Cells the post-write scan checks. Writes are small; a runaway range is not
#: worth minutes of COM round-trips.
MAX_SCAN_CELLS = 2_000

#: Backups kept per workbook. Ten is enough to walk back a bad session and
#: small enough that nobody notices the disk use.
MAX_BACKUPS = 10

#: Where a relative backup root is anchored -- the project root, never the
#: working directory, for the reason ``memory/store.py`` spells out: NEBULA is
#: launched from wherever the user happened to be.
_PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_BACKUP_ROOT = _PROJECT_ROOT / "data" / "excel_backups"

_OPENPYXL_HINT = (
    "Reading Excel files needs openpyxl. Install it with: pip install openpyxl"
)
_XLWINGS_HINT = (
    "Working with the workbook open in Excel needs xlwings. Install it with: "
    "pip install xlwings"
)


class _ExcelUnavailable(RuntimeError):
    """Excel itself could not be reached (not installed, or COM refused)."""


def _fmt(value: Any) -> str:
    """Render one cell for speech.

    Spoken aloud, ``2.0`` and ``datetime.datetime(2026, 9, 1, 0, 0)`` are
    noise; ``2`` and ``2026-09-01`` are the thing the user wrote.
    """
    if value is None:
        return "(empty)"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _fmt_rows(rows: list[tuple[int, list[Any]]]) -> str:
    """Render numbered rows as one spoken clause, never a markdown table."""
    return "; ".join(
        f"row {number} is " + ", ".join(_fmt(cell) for cell in cells)
        for number, cells in rows
    )


def _guess_type(values: list[Any]) -> str:
    """Name what a column holds, from a sample of its cells.

    The model uses this to decide whether ``=SUM`` is even legal on a column
    before it writes the formula, which is cheaper than writing it and reading
    back ``#VALUE!``.
    """
    seen = [value for value in values if value is not None and value != ""]
    if not seen:
        return "empty"
    if any(isinstance(v, str) and v.startswith("=") for v in seen):
        return "formulas"
    if all(isinstance(v, bool) for v in seen):
        return "true/false"
    if all(isinstance(v, (datetime, date)) for v in seen):
        return "dates"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in seen):
        return "numbers"
    return "text"


def _scan(pairs: list[tuple[str, Any]]) -> list[str]:
    """Find Excel error values among (address, value) pairs.

    Compared as text because the same ``#DIV/0!`` arrives as a string from
    openpyxl and, depending on the COM conversion, as displayed text from
    xlwings.
    """
    hits = []
    for address, value in pairs:
        if value is None:
            continue
        text = str(value).strip().upper()
        if text in ERRORS:
            hits.append(f"{text} in {address}")
    return hits


class ExcelSkill(Skill):
    """Read and edit Excel workbooks, on disk or in the running application."""

    name = "ExcelSkill"
    description = "Reads and edits Excel workbooks through openpyxl or live Excel."
    version = "1.0.0"
    enabled = True

    def __init__(
        self,
        container: Any = None,
        backup_root: Path | str | None = None,
    ) -> None:
        super().__init__(container)
        self._backup_root = Path(backup_root) if backup_root else DEFAULT_BACKUP_ROOT
        # Per-workbook, per-session: the point of the backup is the state the
        # workbook was in when NEBULA first touched it. Copying again on the
        # third op would rotate that original out behind our own edits.
        self._backed_up: set[str] = set()

    async def execute(self, task: Task) -> str:
        params = task.parameters or {}
        op = str(params.get("op") or "").strip().lower()

        if task.intent == "excel_read":
            if op not in READ_OPS:
                return REFUSAL
            return self._dispatch_read(op, params)

        if task.intent == "excel_write":
            if op not in WRITE_OPS:
                return REFUSAL
            return self._dispatch_write(op, params)

        raise ValueError(f"ExcelSkill cannot handle intent: {task.intent}")

    # ── backend selection ─────────────────────────────────────────────────

    def _dispatch_read(self, op: str, params: dict[str, Any]) -> str:
        raw_path = params.get("path")

        if not raw_path:
            # No path means "whatever is in front of the user right now",
            # which only the live application can answer.
            return self._read_live(op, None, params)

        path = self._resolve(raw_path)
        if not path.exists():
            return f"I can't find that workbook: {path}"
        if path.is_dir():
            return f"{path} is a folder, not a workbook."
        if path.suffix.lower() not in (".xlsx", ".xlsm"):
            return (
                f"{path.name} isn't an xlsx workbook, so I can't open it as one."
            )

        try:
            return self._read_file(op, path, params)
        except ImportError:
            return _OPENPYXL_HINT
        except PermissionError:
            # Excel has the file locked. The live copy is the authoritative
            # one anyway, so read that instead of reporting a lock.
            return self._read_live(op, path, params)

    def _dispatch_write(self, op: str, params: dict[str, Any]) -> str:
        raw_path = params.get("path")

        if not raw_path:
            return self._write_live(op, None, params)

        path = self._resolve(raw_path)
        if not path.exists():
            return f"I can't find that workbook: {path}"
        if path.is_dir():
            return f"{path} is a folder, not a workbook."
        if path.suffix.lower() == ".xlsm":
            # openpyxl drops the macro project unless the workbook is loaded
            # with keep_vba, and a silently de-macro'd workbook is a bug the
            # user finds out about a week later.
            return (
                f"{path.name} is a macro workbook. Saving it would strip the "
                f"macros, so I won't write to it."
            )
        if path.suffix.lower() != ".xlsx":
            return f"{path.name} isn't an xlsx workbook, so I can't write to it."

        ok, note = self._ensure_backup(path)
        if not ok:
            return note

        try:
            return self._write_file(op, path, params, backup_note=note)
        except ImportError:
            return _OPENPYXL_HINT
        except PermissionError:
            # The classic case: the user is looking at the file in Excel. The
            # edit is not impossible, it just has to go through Excel.
            logger.info(f"{path.name} is locked by Excel; retrying through xlwings")
            return self._write_live(op, path, params, backup_note=note)

    # ── backups ───────────────────────────────────────────────────────────

    def _ensure_backup(self, path: Path) -> tuple[bool, str]:
        """Copy the workbook aside once per session, and prune old copies.

        Returns (ok, sentence). A failed backup refuses the write outright,
        the same rule ``code_skill._write`` follows: an edit nobody can undo
        does not happen.
        """
        key = str(path).lower()
        if key in self._backed_up:
            return True, ""

        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        try:
            self._backup_root.mkdir(parents=True, exist_ok=True)
            target = self._backup_root / f"{path.stem}_{stamp}{path.suffix}"
            # Two backups inside one second only happens under test or on a
            # very fast agent loop; suffixing beats overwriting the earlier one.
            counter = 2
            while target.exists():
                target = self._backup_root / f"{path.stem}_{stamp}-{counter}{path.suffix}"
                counter += 1
            shutil.copy2(path, target)
        except OSError as exc:
            return False, (
                f"I couldn't back up {path.name} ({exc}), so I haven't changed it."
            )

        self._backed_up.add(key)
        self._prune(path)
        logger.info(f"Backed up {path.name} to {target}")
        return True, f" I backed the workbook up first, as {target.name}."

    def _prune(self, path: Path) -> None:
        """Keep the newest MAX_BACKUPS copies *of this workbook*.

        Per workbook rather than per folder: ten sessions against a budget
        should not evict the only copy of a different spreadsheet.

        Which is exactly what a ``{stem}_*{suffix}`` glob did. For
        ``report.xlsx`` it also matched ``report_2025_<stamp>.xlsx`` -- the
        backups of a *different* workbook, ``report_2025.xlsx`` -- and since
        those sort earlier they were the first thing deleted. NEBULA silently
        destroyed the undo history of a file the user never asked it to
        touch. The name has to be matched exactly, timestamp and all.

        Ordering is by modification time, not by name: the same-second
        collision suffix makes names sort wrongly (``-10`` before ``-2``
        before no suffix at all), which evicted newer copies and kept older
        ones.
        """
        pattern = re.compile(
            rf"^{re.escape(path.stem)}_\d{{8}}-\d{{6}}(-\d+)?{re.escape(path.suffix)}$",
            re.IGNORECASE,
        )
        try:
            copies = sorted(
                (
                    item
                    for item in self._backup_root.glob(f"{path.stem}_*{path.suffix}")
                    if pattern.match(item.name)
                ),
                key=lambda item: item.stat().st_mtime,
            )
            for stale in copies[:-MAX_BACKUPS]:
                stale.unlink(missing_ok=True)
        except OSError as exc:  # pragma: no cover - pruning is best effort
            logger.warning(f"Could not prune Excel backups: {exc}")

    # ── openpyxl: reads ───────────────────────────────────────────────────

    def _read_file(self, op: str, path: Path, params: dict[str, Any]) -> str:
        from openpyxl import load_workbook

        # Not read_only=True: it reports no dimensions for workbooks written
        # by tools that omit them, and everything here is capped anyway, so
        # correctness wins over the load time.
        values_wb = load_workbook(path, data_only=True)

        try:
            if op == "sheets":
                names = values_wb.sheetnames
                return (
                    f"{path.name} has {len(names)} "
                    f"{'sheet' if len(names) == 1 else 'sheets'}: "
                    f"{', '.join(names)}."
                )

            sheet_name = params.get("sheet") or values_wb.sheetnames[0]
            if sheet_name not in values_wb.sheetnames:
                return (
                    f"There's no sheet called '{sheet_name}' in {path.name}. "
                    f"The sheets are: {', '.join(values_wb.sheetnames)}."
                )

            if op == "shape":
                # A second load with formulas intact: data_only shows what
                # Excel last calculated, which is None for anything openpyxl
                # itself wrote, so the formula view is what tells the model a
                # column is computed rather than typed.
                formulas_wb = load_workbook(path, data_only=False)
                try:
                    return self._shape(
                        path.name,
                        values_wb,
                        values_wb[sheet_name],
                        formulas_wb[sheet_name],
                    )
                finally:
                    formulas_wb.close()

            return self._read_range(
                path.name, values_wb[sheet_name], params.get("range")
            )
        finally:
            values_wb.close()

    def _shape(self, label: str, workbook: Any, values_ws: Any, formulas_ws: Any) -> str:
        max_row = values_ws.max_row or 0
        max_col = values_ws.max_column or 0

        if max_row == 0 or max_col == 0:
            return f"Sheet '{values_ws.title}' in {label} is empty."

        header = self._row(values_ws, 1, max_col)
        if all(cell is None for cell in header):
            header = self._row(formulas_ws, 1, max_col)

        data_rows = [
            (number, self._row(values_ws, number, max_col, fallback=formulas_ws))
            for number in range(2, min(max_row, 1 + SHAPE_ROWS) + 1)
        ]

        samples = [
            self._row(formulas_ws, number, max_col)
            for number in range(2, min(max_row, 1 + TYPE_SAMPLE_ROWS) + 1)
        ]
        types = []
        for index in range(max_col):
            column = [row[index] for row in samples if index < len(row)]
            name = _fmt(header[index]) if index < len(header) else ""
            label_for_column = name if name and name != "(empty)" else f"column {index + 1}"
            types.append(f"{label_for_column} holds {_guess_type(column)}")

        parts = [
            f"{label} has {len(workbook.sheetnames)} "
            f"{'sheet' if len(workbook.sheetnames) == 1 else 'sheets'}: "
            f"{', '.join(workbook.sheetnames)}.",
            f"Sheet '{values_ws.title}' uses {values_ws.dimensions} - "
            f"{max_col} {'column' if max_col == 1 else 'columns'} by "
            f"{max_row} {'row' if max_row == 1 else 'rows'}.",
            f"The header row is: {', '.join(_fmt(cell) for cell in header)}.",
            f"Column types: {'; '.join(types)}.",
        ]

        if data_rows:
            total_data = max(max_row - 1, 0)
            shown = len(data_rows)
            lead = (
                f"First {shown} of {total_data} data rows"
                if total_data > shown
                else f"All {shown} data rows"
            )
            parts.append(f"{lead}: {_fmt_rows(data_rows)}.")
        else:
            parts.append("There are no data rows below the header.")

        return " ".join(parts)

    @staticmethod
    def _row(worksheet: Any, number: int, max_col: int, fallback: Any = None) -> list[Any]:
        cells = [
            worksheet.cell(row=number, column=column).value
            for column in range(1, max_col + 1)
        ]
        if fallback is not None:
            # A formula openpyxl wrote has no cached value; showing the
            # formula beats showing an empty cell that is not empty.
            cells = [
                cell
                if cell is not None
                else fallback.cell(row=number, column=column).value
                for column, cell in enumerate(cells, start=1)
            ]
        return cells

    def _read_range(self, label: str, worksheet: Any, raw_range: Any) -> str:
        if not raw_range:
            return "I need a range to read, like A1:D20."

        bounds = self._bounds(str(raw_range))
        if bounds is None:
            return (
                f"I couldn't make sense of the range '{raw_range}'. Use A1 "
                f"style, like A1:D20."
            )
        min_col, min_row, max_col, max_row = bounds

        rows: list[tuple[int, list[Any]]] = []
        cells_read = 0
        truncated = False
        for number in range(min_row, max_row + 1):
            if cells_read >= MAX_RANGE_CELLS:
                truncated = True
                break
            values = [
                worksheet.cell(row=number, column=column).value
                for column in range(min_col, max_col + 1)
            ]
            cells_read += len(values)
            rows.append((number, values))

        if not rows:
            return f"{raw_range} on '{worksheet.title}' in {label} is empty."

        note = (
            f" That's the first {len(rows)} rows of the range; ask for a "
            f"smaller one to see the rest."
            if truncated
            else ""
        )
        return (
            f"{raw_range} on '{worksheet.title}' in {label}: "
            f"{_fmt_rows(rows)}.{note}"
        )

    # ── openpyxl: writes ──────────────────────────────────────────────────

    def _write_file(
        self, op: str, path: Path, params: dict[str, Any], backup_note: str = ""
    ) -> str:
        from openpyxl import load_workbook

        # data_only=False is not a default, it is the whole point: loading a
        # workbook with data_only=True and saving it replaces every formula
        # with its last cached value. That turns one edit into a workbook that
        # no longer recalculates -- silent, total, and unnoticed for months.
        workbook = load_workbook(path, data_only=False)

        try:
            if op == "add_sheet":
                name = str(params.get("name") or "").strip()
                if not name:
                    return "I need a name for the new sheet."
                if name in workbook.sheetnames:
                    return f"There's already a sheet called '{name}' in {path.name}."
                workbook.create_sheet(title=name)
                workbook.save(path)
                return (
                    f"Added a sheet called '{name}' to {path.name}. The "
                    f"workbook now has {len(workbook.sheetnames)} sheets."
                    f"{backup_note}"
                )

            sheet_name = params.get("sheet") or workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                return (
                    f"There's no sheet called '{sheet_name}' in {path.name}. "
                    f"The sheets are: {', '.join(workbook.sheetnames)}."
                )
            worksheet = workbook[sheet_name]

            if op == "set_values":
                outcome = self._set_values_file(worksheet, params)
            elif op == "set_formula":
                outcome = self._set_formula_file(worksheet, params)
            else:
                outcome = self._format_bold_file(worksheet, params)

            if isinstance(outcome, str):
                return outcome
            written, sentence = outcome

            workbook.save(path)
        finally:
            workbook.close()

        errors = self._scan_file(path, sheet_name, written)
        return self._write_reply(
            f"{sentence} on '{sheet_name}' in {path.name}",
            errors,
            recalculated=False,
            backup_note=backup_note,
        )

    def _set_values_file(
        self, worksheet: Any, params: dict[str, Any]
    ) -> str | tuple[str, str]:
        values = self._values(params.get("values"))
        if isinstance(values, str):
            return values

        raw_range = params.get("range")
        if not raw_range:
            return "I need the target range for set_values, like A2:D10."
        bounds = self._bounds(str(raw_range))
        if bounds is None:
            return (
                f"I couldn't make sense of the range '{raw_range}'. Use A1 "
                f"style, like A2:D10."
            )
        min_col, min_row, _, _ = bounds

        count = 0
        for row_offset, row in enumerate(values):
            for col_offset, value in enumerate(row):
                worksheet.cell(
                    row=min_row + row_offset, column=min_col + col_offset, value=value
                )
                count += 1

        written = self._span(
            min_col, min_row, len(values), max(len(row) for row in values)
        )
        note = ""
        if written != str(raw_range).upper().replace("$", ""):
            # §2.5's second half: say where the values actually landed, so a
            # model that sent a 3-row block for a 2-row range finds out.
            note = f" (you asked for {raw_range})"
        return written, f"Wrote {count} values to {written}{note}"

    def _set_formula_file(
        self, worksheet: Any, params: dict[str, Any]
    ) -> str | tuple[str, str]:
        formula = str(params.get("formula") or "").strip()
        if not formula:
            return "I need the formula to write, like =B2*0.18."
        if not formula.startswith("="):
            formula = "=" + formula

        raw_range = params.get("range")
        if not raw_range:
            return "I need the target range for set_formula, like C2:C20."
        bounds = self._bounds(str(raw_range))
        if bounds is None:
            return (
                f"I couldn't make sense of the range '{raw_range}'. Use A1 "
                f"style, like C2:C20."
            )
        min_col, min_row, max_col, max_row = bounds

        from openpyxl.formula.translate import Translator
        from openpyxl.utils import get_column_letter

        anchor = f"{get_column_letter(min_col)}{min_row}"
        worksheet[anchor] = formula
        count = 1
        # Filling a column by hand would write =B2*0.18 into every row. Excel
        # shifts the references when you drag a formula down; Translator is
        # the same operation, so the model can say "this column" once.
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                if row == min_row and column == min_col:
                    continue
                destination = f"{get_column_letter(column)}{row}"
                worksheet[destination] = Translator(
                    formula, origin=anchor
                ).translate_formula(destination)
                count += 1

        written = self._span(min_col, min_row, max_row - min_row + 1, max_col - min_col + 1)
        return written, f"Put {formula} into {count} cells across {written}"

    def _format_bold_file(
        self, worksheet: Any, params: dict[str, Any]
    ) -> str | tuple[str, str]:
        raw_range = params.get("range")
        if not raw_range:
            # "bold the header" is the request this op exists for, and the
            # header is row 1 in every sheet shape() can describe.
            bounds = (1, 1, worksheet.max_column or 1, 1)
        else:
            found = self._bounds(str(raw_range))
            if found is None:
                return (
                    f"I couldn't make sense of the range '{raw_range}'. Use A1 "
                    f"style, like A1:D1."
                )
            bounds = found

        min_col, min_row, max_col, max_row = bounds
        count = 0
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = worksheet.cell(row=row, column=column)
                # Copy the existing font rather than building a new one: a new
                # Font() would reset the size and typeface to the defaults.
                font = _copy_style(cell.font)
                font.bold = True
                cell.font = font
                count += 1

        written = self._span(min_col, min_row, max_row - min_row + 1, max_col - min_col + 1)
        return written, f"Made {count} cells bold across {written}"

    def _scan_file(self, path: Path, sheet_name: str, written: str) -> list[str]:
        """Re-read what was just written and look for Excel error values.

        Reads the saved file rather than the in-memory workbook: the question
        is what is on disk now, not what we believed we wrote.
        """
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        bounds = self._bounds(written)
        if bounds is None:  # pragma: no cover - written is built by us
            return []
        min_col, min_row, max_col, max_row = bounds

        workbook = load_workbook(path, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            pairs: list[tuple[str, Any]] = []
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    if len(pairs) >= MAX_SCAN_CELLS:
                        break
                    address = f"{get_column_letter(column)}{row}"
                    pairs.append(
                        (address, worksheet.cell(row=row, column=column).value)
                    )
            return _scan(pairs)
        finally:
            workbook.close()

    # ── xlwings: the live application ─────────────────────────────────────

    def _read_live(self, op: str, path: Path | None, params: dict[str, Any]) -> str:
        try:
            book = self._book(path)
        except ImportError:
            return _XLWINGS_HINT
        except _ExcelUnavailable as exc:
            return str(exc)

        try:
            label = book.name
            if op == "sheets":
                names = [sheet.name for sheet in book.sheets]
                return (
                    f"{label} has {len(names)} "
                    f"{'sheet' if len(names) == 1 else 'sheets'}: "
                    f"{', '.join(names)}."
                )

            sheet = self._sheet(book, params.get("sheet"))
            if sheet is None:
                names = ", ".join(s.name for s in book.sheets)
                return (
                    f"There's no sheet called '{params.get('sheet')}' in "
                    f"{label}. The sheets are: {names}."
                )

            if op == "shape":
                return self._shape_live(label, book, sheet)
            return self._read_range_live(label, sheet, params.get("range"))
        except _ExcelUnavailable as exc:
            return str(exc)
        except Exception as exc:  # COM raises a zoo of types; none are useful aloud
            logger.warning(f"Excel read through xlwings failed: {exc}")
            return f"Excel wouldn't let me read that: {exc}"

    def _shape_live(self, label: str, book: Any, sheet: Any) -> str:
        used = sheet.used_range
        grid = used.value
        if grid is None:
            return f"Sheet '{sheet.name}' in {label} is empty."
        if not isinstance(grid, list):
            grid = [[grid]]
        elif grid and not isinstance(grid[0], list):
            grid = [grid]

        header = grid[0]
        data = grid[1:]
        shown = [(index + 2, row) for index, row in enumerate(data[:SHAPE_ROWS])]
        samples = data[:TYPE_SAMPLE_ROWS]

        types = []
        for index in range(len(header)):
            column = [row[index] for row in samples if index < len(row)]
            name = _fmt(header[index])
            label_for_column = name if name != "(empty)" else f"column {index + 1}"
            types.append(f"{label_for_column} holds {_guess_type(column)}")

        parts = [
            f"{label} has {len(book.sheets)} sheets: "
            f"{', '.join(s.name for s in book.sheets)}.",
            f"Sheet '{sheet.name}' uses {used.address.replace('$', '')} - "
            f"{len(header)} columns by {len(grid)} rows.",
            f"The header row is: {', '.join(_fmt(cell) for cell in header)}.",
            f"Column types: {'; '.join(types)}.",
        ]
        if shown:
            lead = (
                f"First {len(shown)} of {len(data)} data rows"
                if len(data) > len(shown)
                else f"All {len(shown)} data rows"
            )
            parts.append(f"{lead}: {_fmt_rows(shown)}.")
        else:
            parts.append("There are no data rows below the header.")
        return " ".join(parts)

    def _read_range_live(self, label: str, sheet: Any, raw_range: Any) -> str:
        if not raw_range:
            return "I need a range to read, like A1:D20."

        bounds = self._bounds(str(raw_range))
        if bounds is None:
            return (
                f"I couldn't make sense of the range '{raw_range}'. Use A1 "
                f"style, like A1:D20."
            )
        min_col, min_row, max_col, max_row = bounds

        grid = sheet.range(str(raw_range)).value
        if not isinstance(grid, list):
            grid = [[grid]]
        elif grid and not isinstance(grid[0], list):
            grid = [grid] if max_row == min_row else [[cell] for cell in grid]

        rows: list[tuple[int, list[Any]]] = []
        cells = 0
        truncated = False
        for offset, row in enumerate(grid):
            if cells >= MAX_RANGE_CELLS:
                truncated = True
                break
            rows.append((min_row + offset, list(row)))
            cells += len(row)

        note = (
            f" That's the first {len(rows)} rows of the range; ask for a "
            f"smaller one to see the rest."
            if truncated
            else ""
        )
        return (
            f"{raw_range} on '{sheet.name}' in {label}: {_fmt_rows(rows)}.{note}"
        )

    def _write_live(
        self,
        op: str,
        path: Path | None,
        params: dict[str, Any],
        backup_note: str = "",
    ) -> str:
        try:
            book = self._book(path)
        except ImportError:
            return _XLWINGS_HINT
        except _ExcelUnavailable as exc:
            return str(exc)

        try:
            label = book.name

            if not backup_note:
                # The live path reaches here without a backup when no path was
                # given. The workbook still has a file on disk to copy unless
                # it has never been saved, and Excel's lock permits reading.
                on_disk = self._book_path(book)
                if on_disk is not None:
                    ok, backup_note = self._ensure_backup(on_disk)
                    if not ok:
                        return backup_note
                else:
                    backup_note = (
                        " This workbook has never been saved, so there was "
                        "nothing to back up."
                    )

            if op == "add_sheet":
                name = str(params.get("name") or "").strip()
                if not name:
                    return "I need a name for the new sheet."
                if any(sheet.name == name for sheet in book.sheets):
                    return f"There's already a sheet called '{name}' in {label}."
                book.sheets.add(name, after=book.sheets[-1])
                book.save()
                return (
                    f"Added a sheet called '{name}' to {label}. The workbook "
                    f"now has {len(book.sheets)} sheets.{backup_note}"
                )

            sheet = self._sheet(book, params.get("sheet"))
            if sheet is None:
                names = ", ".join(s.name for s in book.sheets)
                return (
                    f"There's no sheet called '{params.get('sheet')}' in "
                    f"{label}. The sheets are: {names}."
                )

            raw_range = params.get("range")
            if op in ("set_values", "set_formula") and not raw_range:
                return f"I need the target range for {op}, like A2:D10."

            if op == "set_values":
                values = self._values(params.get("values"))
                if isinstance(values, str):
                    return values
                bounds = self._bounds(str(raw_range))
                if bounds is None:
                    return (
                        f"I couldn't make sense of the range '{raw_range}'. "
                        f"Use A1 style, like A2:D10."
                    )
                min_col, min_row, _, _ = bounds
                written = self._span(
                    min_col, min_row, len(values), max(len(row) for row in values)
                )
                sheet.range(written).value = values
                sentence = f"Wrote {sum(len(row) for row in values)} values to {written}"

            elif op == "set_formula":
                formula = str(params.get("formula") or "").strip()
                if not formula:
                    return "I need the formula to write, like =B2*0.18."
                if not formula.startswith("="):
                    formula = "=" + formula
                written = str(raw_range).upper().replace("$", "")
                # Assigning one formula to a multi-cell range is Excel's own
                # fill: it shifts the relative references per cell for us.
                sheet.range(written).formula = formula
                sentence = f"Put {formula} across {written}"

            else:  # format_bold
                # No range means "the header", which live Excel can find for
                # itself: A1 expanded rightwards is the contiguous header row.
                written = (
                    str(raw_range).replace("$", "").upper()
                    if raw_range
                    else sheet.range("A1").expand("right").address.replace("$", "")
                )
                target = sheet.range(written)
                try:
                    target.font.bold = True
                except AttributeError:  # pragma: no cover - older xlwings
                    target.api.Font.Bold = True
                sentence = f"Made {written} bold"

            book.save()
            errors = self._scan_live(sheet, written)
            return self._write_reply(
                f"{sentence} on '{sheet.name}' in {label}",
                errors,
                recalculated=True,
                backup_note=backup_note,
            )
        except _ExcelUnavailable as exc:
            return str(exc)
        except Exception as exc:
            logger.warning(f"Excel write through xlwings failed: {exc}")
            return f"Excel wouldn't let me make that change: {exc}"

    def _scan_live(self, sheet: Any, written: str) -> list[str]:
        """Scan the live range for error values.

        This is the scan that actually means something: Excel has recalculated
        by the time we read, so a formula that broke says so now rather than
        the next time somebody opens the file.
        """
        from openpyxl.utils import get_column_letter

        bounds = self._bounds(written)
        if bounds is None:  # pragma: no cover
            return []
        min_col, min_row, max_col, max_row = bounds

        pairs: list[tuple[str, Any]] = []
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                if len(pairs) >= MAX_SCAN_CELLS:
                    break
                address = f"{get_column_letter(column)}{row}"
                cell = sheet.range(address)
                value = cell.value
                if value is None or str(value).strip().upper() not in ERRORS:
                    # COM hands error values back in several shapes depending
                    # on the converter; the displayed text is the one that
                    # always reads "#DIV/0!".
                    try:
                        value = cell.api.Text
                    except Exception:  # pragma: no cover - COM detail
                        pass
                pairs.append((address, value))
        return _scan(pairs)

    def _book(self, path: Path | None) -> Any:
        """Find the live workbook, or say plainly why we cannot."""
        import xlwings as xw

        try:
            if path is None:
                book = xw.books.active
                if book is None:
                    raise _ExcelUnavailable(
                        "There's no workbook open in Excel, so tell me the file "
                        "path and I'll open it from disk."
                    )
                return book

            for candidate in xw.books:
                try:
                    if Path(candidate.fullname).resolve() == path.resolve():
                        return candidate
                except Exception:  # pragma: no cover - unsaved books have no path
                    continue
            return xw.Book(str(path))
        except _ExcelUnavailable:
            raise
        except Exception as exc:
            # No Excel, no COM, or Excel busy in a modal dialog. All of these
            # are one sentence to the user, never a traceback.
            logger.warning(f"Could not reach Excel: {exc}")
            raise _ExcelUnavailable(
                "I couldn't reach Excel on this machine - it needs Excel "
                "installed and not stuck in a dialog. Give me the file path "
                "and I can work on the closed file instead."
            ) from exc

    @staticmethod
    def _book_path(book: Any) -> Path | None:
        try:
            full = book.fullname
        except Exception:  # pragma: no cover - COM detail
            return None
        candidate = Path(str(full))
        return candidate if candidate.exists() else None

    @staticmethod
    def _sheet(book: Any, name: Any) -> Any:
        if not name:
            return book.sheets.active
        for sheet in book.sheets:
            if sheet.name == str(name):
                return sheet
        return None

    # ── shared helpers ────────────────────────────────────────────────────

    def _write_reply(
        self,
        sentence: str,
        errors: list[str],
        *,
        recalculated: bool,
        backup_note: str,
    ) -> str:
        if errors:
            found = " and ".join(errors)
            tail = (
                "Those are the values Excel calculated."
                if recalculated
                else "Excel hasn't recalculated the file yet, so there may be more."
            )
            return f"{sentence}, but the range now contains {found}. {tail}{backup_note}"

        clean = (
            "I read the range back and Excel shows no error values there."
            if recalculated
            else (
                "I read the range back and found no error values, though these "
                "are the stored values - I edited the file on disk, so Excel "
                "hasn't recalculated it."
            )
        )
        return f"{sentence}. {clean}{backup_note}"

    @staticmethod
    def _values(raw: Any) -> list[list[Any]] | str:
        """Coerce the model's ``values`` argument into a 2-D list.

        Tool arguments arrive as JSON, but a small model sends the array as a
        string about as often as it sends a list, and a single row without the
        outer brackets more often than either.
        """
        if raw is None or raw == "":
            return "I need the values to write, as a 2-D array like [[1, 2], [3, 4]]."

        if isinstance(raw, str):
            try:
                raw = json.loads(raw)
            except json.JSONDecodeError:
                return (
                    "I couldn't read those values. Send them as a JSON array of "
                    "rows, like [[1, 2], [3, 4]]."
                )

        if not isinstance(raw, list) or not raw:
            return "I need the values to write, as a 2-D array like [[1, 2], [3, 4]]."

        if not any(isinstance(row, list) for row in raw):
            raw = [raw]

        rows = [row if isinstance(row, list) else [row] for row in raw]
        return rows

    @staticmethod
    def _bounds(raw_range: str) -> tuple[int, int, int, int] | None:
        from openpyxl.utils import range_boundaries

        try:
            min_col, min_row, max_col, max_row = range_boundaries(
                str(raw_range).replace("$", "").strip().upper()
            )
        except Exception:
            return None
        if None in (min_col, min_row, max_col, max_row):
            # Whole-column ("A:A") and whole-row ranges come back half-open.
            return None
        return int(min_col), int(min_row), int(max_col), int(max_row)

    @staticmethod
    def _span(min_col: int, min_row: int, rows: int, cols: int) -> str:
        from openpyxl.utils import get_column_letter

        start = f"{get_column_letter(min_col)}{min_row}"
        end = f"{get_column_letter(min_col + cols - 1)}{min_row + rows - 1}"
        return start if start == end else f"{start}:{end}"

    @staticmethod
    def _resolve(raw_path: Any) -> Path:
        """Expand ``~`` and variables, the same way ``code_skill`` does.

        Model-supplied paths are untrusted input; expanding them here and
        nowhere else keeps one rule instead of four.
        """
        expanded = os.path.expandvars(os.path.expanduser(str(raw_path).strip()))
        return Path(expanded).expanduser().absolute()
